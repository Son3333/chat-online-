# -*- coding: utf-8 -*-
"""
Script gộp toàn bộ các chương 1 -> 5 thành cuốn toàn văn Báo cáo Đồ án Chuyên ngành
Chuẩn Nhóm 21 - Trường Đại học Công nghệ Đông Á
"""

import os
import sys
import shutil
from copy import deepcopy
import docx
from docx.oxml.section import CT_SectPr

sys.stdout.reconfigure(encoding='utf-8')

def build_merged_report():
    print("[*] Đang khởi tạo bản gộp toàn văn từ Chương 1...")
    master_doc = docx.Document('BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH - CHƯƠNG 1 - NHÓM 21.docx')
    
    chapter_files = [
        ('Chương 2', 'BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH - CHƯƠNG 2 - NHÓM 21.docx'),
        ('Chương 3', 'BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH - CHƯƠNG 3 - NHÓM 21.docx'),
        ('Chương 4', 'BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH - CHƯƠNG 4 - NHÓM 21.docx'),
        ('Chương 5 & Kết Luận', 'BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH - CHƯƠNG 5 - KẾT LUẬN - NHÓM 21.docx'),
    ]
    
    for ch_name, filename in chapter_files:
        print(f"[*] Đang ghép nối {ch_name} ({filename})...")
        ch_doc = docx.Document(filename)
        
        # Thêm ngắt trang trước khi vào chương mới
        master_doc.add_page_break()
        
        # Bỏ qua p0, p1, tbl0, p2 (phần hộp tiêu đề trường/nhóm đã có ở đầu toàn văn)
        # Bắt đầu lấy từ phần tử p3 (Tiêu đề chương) trở đi
        body_children = list(ch_doc.element.body)
        
        # Xác định vị trí bắt đầu nội dung chương (tìm phần tử chứa 'CHƯƠNG')
        start_idx = 0
        for idx, child in enumerate(body_children):
            if child.tag.endswith('p'):
                text = "".join(child.itertext())
                if "CHƯƠNG" in text.upper():
                    start_idx = idx
                    break
        
        for idx in range(start_idx, len(body_children)):
            child = body_children[idx]
            if isinstance(child, CT_SectPr):
                continue
            master_doc.element.body.append(deepcopy(child))
            
    out_filename = "BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH - TOÀN VĂN ĐỒ ÁN - NHÓM 21.docx"
    master_doc.save(out_filename)
    shutil.copyfile(out_filename, os.path.join("docs", out_filename))
    
    total_p = len(master_doc.paragraphs)
    total_t = len(master_doc.tables)
    words = sum(len(p.text.split()) for p in master_doc.paragraphs)
    
    print("\n" + "="*60)
    print(f"[SUCCESS] ĐÃ XUẤT THÀNH CÔNG BÁO CÁO TOÀN VĂN:")
    print(f"  - Tập tin: {out_filename}")
    print(f"  - Thư mục docs/: docs/{out_filename}")
    print(f"  - Tổng số đoạn văn: {total_p}")
    print(f"  - Tổng số bảng biểu: {total_t}")
    print(f"  - Tổng số từ: ~{words} từ")
    print("="*60)

if __name__ == '__main__':
    build_merged_report()
